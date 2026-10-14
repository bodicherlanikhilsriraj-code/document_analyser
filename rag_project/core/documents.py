"""Safe text extraction for supported student study materials."""

from __future__ import annotations

import hashlib
import io
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path

from .errors import DocumentError

SUPPORTED_EXTENSIONS = {".txt", ".pdf", ".docx"}
MAX_EXTRACTED_CHARACTERS = 20_000_000
MAX_DOCX_EXPANDED_BYTES = 100 * 1024 * 1024


@dataclass(frozen=True)
class ExtractedPage:
    page_number: int | None
    text: str


def document_hash(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def read_stored_document(path: Path, upload_dir: Path, max_bytes: int) -> bytes:
    if path.resolve().parent != upload_dir.resolve():
        raise DocumentError("The local source path is outside the StudyPilot upload folder.")
    try:
        if path.stat().st_size > max_bytes:
            raise DocumentError("The local source file exceeds the configured upload limit.")
        content = path.read_bytes()
    except OSError as exc:
        raise DocumentError(
            "The local source file is missing or unreadable. Upload the document again."
        ) from exc
    if len(content) > max_bytes:
        raise DocumentError("The local source file exceeds the configured upload limit.")
    return content


def _clean_text(text: str) -> str:
    text = text.replace("\x00", "")
    return re.sub(r"[\x01-\x08\x0b\x0c\x0e-\x1f]", "", text).strip()


def extract_document(
    filename: str, content: bytes, max_bytes: int = 25 * 1024 * 1024
) -> list[ExtractedPage]:
    """Extract non-empty text while retaining page metadata when available."""
    if not filename or not filename.strip():
        raise DocumentError("Choose a file with a valid filename.")
    suffix = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if suffix not in SUPPORTED_EXTENSIONS:
        raise DocumentError("Unsupported file type. Upload a PDF, DOCX, or TXT file.")
    if not content:
        raise DocumentError("The selected file is empty.")
    if len(content) > max_bytes:
        raise DocumentError(f"The file exceeds the {max_bytes // (1024 * 1024)} MB limit.")

    pages: list[ExtractedPage]
    try:
        if suffix == ".txt":
            text = None
            chosen_encoding = ""
            for encoding in ("utf-8-sig", "utf-16", "cp1252"):
                try:
                    text = content.decode(encoding)
                    chosen_encoding = encoding
                    break
                except UnicodeDecodeError:
                    continue
            if text is None:
                raise DocumentError("This text file encoding is not supported.")
            if "\x00" in text and chosen_encoding not in {"utf-16"}:
                raise DocumentError("This file does not appear to contain readable text.")
            control_count = sum(
                ord(character) < 32 and character not in "\t\r\n\f" for character in text
            )
            if text and control_count / len(text) > 0.01:
                raise DocumentError("This file does not appear to contain readable text.")
            pages = [ExtractedPage(None, _clean_text(text))]
        elif suffix == ".pdf":
            import pymupdf

            with pymupdf.open(stream=content, filetype="pdf") as pdf:
                if pdf.is_encrypted:
                    raise DocumentError("This PDF is password protected.")
                if len(pdf) > 20_000:
                    raise DocumentError("This PDF exceeds the 20,000-page processing limit.")
                pages = []
                total_characters = 0
                for index, page in enumerate(pdf):
                    text = _clean_text(page.get_text("text"))
                    total_characters += len(text)
                    if total_characters > MAX_EXTRACTED_CHARACTERS:
                        raise DocumentError("The extracted PDF text exceeds the safe processing limit.")
                    pages.append(ExtractedPage(index + 1, text))
        else:
            from docx import Document

            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                entries = archive.infolist()
                if len(entries) > 10_000 or sum(item.file_size for item in entries) > MAX_DOCX_EXPANDED_BYTES:
                    raise DocumentError("The DOCX expands beyond the safe processing limit.")
            doc = Document(io.BytesIO(content))
            blocks = [paragraph.text.strip() for paragraph in doc.paragraphs if paragraph.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        blocks.append(" | ".join(cells))
            text = _clean_text("\n".join(blocks))
            if len(text) > MAX_EXTRACTED_CHARACTERS:
                raise DocumentError("The extracted DOCX text exceeds the safe processing limit.")
            pages = [ExtractedPage(None, text)]
    except DocumentError:
        raise
    except ImportError as exc:
        package = "PyMuPDF" if suffix == ".pdf" else "python-docx"
        raise DocumentError(
            f"The {package} parser is not installed. Install the project requirements and retry."
        ) from exc
    except Exception as exc:
        raise DocumentError(
            "The file could not be read. Check that it is a valid, uncorrupted document."
        ) from exc

    pages = [page for page in pages if page.text]
    if not pages:
        if suffix == ".pdf":
            raise DocumentError(
                "No selectable text was found in this PDF. It may be scanned; OCR is not "
                "included, so use a text-based PDF or OCR the file before uploading."
            )
        raise DocumentError("No readable text was found in this document.")
    return pages
