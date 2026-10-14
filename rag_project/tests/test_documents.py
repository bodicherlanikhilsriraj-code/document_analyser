import importlib.util
import tempfile
import unittest
from pathlib import Path

from core.documents import document_hash, extract_document, read_stored_document
from core.errors import DocumentError


class DocumentTests(unittest.TestCase):
    def test_extracts_utf8_text(self):
        pages = extract_document("lecture.txt", "Study notes".encode())
        self.assertEqual(pages[0].text, "Study notes")
        self.assertIsNone(pages[0].page_number)

    def test_rejects_empty_unsupported_and_large_files(self):
        with self.assertRaises(DocumentError):
            extract_document("empty.txt", b"")
        with self.assertRaises(DocumentError):
            extract_document("slides.pptx", b"data")
        with self.assertRaises(DocumentError):
            extract_document("large.txt", b"1234", max_bytes=3)
        with self.assertRaises(DocumentError):
            extract_document("blank.txt", b" \n\t ")
        with self.assertRaises(DocumentError):
            extract_document("binary.txt", b"\x00\x01\x02")

    def test_hash_is_deterministic(self):
        self.assertEqual(document_hash(b"same"), document_hash(b"same"))
        self.assertNotEqual(document_hash(b"same"), document_hash(b"other"))

    def test_stored_document_reads_are_bounded_to_upload_directory(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            root = Path(temporary_directory)
            upload_dir = root / "uploads"
            upload_dir.mkdir()
            source = upload_dir / "study.txt"
            source.write_bytes(b"notes")
            self.assertEqual(read_stored_document(source, upload_dir, 10), b"notes")
            with self.assertRaises(DocumentError):
                read_stored_document(source, upload_dir, 2)
            outside = root / "private.txt"
            outside.write_bytes(b"private")
            with self.assertRaises(DocumentError):
                read_stored_document(outside, upload_dir, 100)

    @unittest.skipUnless(importlib.util.find_spec("pymupdf"), "PyMuPDF is not installed")
    def test_extracts_pdf_with_page_numbers(self):
        import pymupdf

        pdf = pymupdf.open()
        page = pdf.new_page()
        page.insert_text((72, 72), "A real PDF passage")
        content = pdf.tobytes()
        pdf.close()
        pages = extract_document("notes.pdf", content)
        self.assertEqual(pages[0].page_number, 1)
        self.assertIn("real PDF passage", pages[0].text)

    @unittest.skipUnless(importlib.util.find_spec("pymupdf"), "PyMuPDF is not installed")
    def test_rejects_corrupt_and_textless_pdfs(self):
        import pymupdf

        with self.assertRaises(DocumentError):
            extract_document("corrupt.pdf", b"not a PDF")
        pdf = pymupdf.open()
        pdf.new_page()
        content = pdf.tobytes()
        pdf.close()
        with self.assertRaisesRegex(DocumentError, "OCR"):
            extract_document("scanned.pdf", content)

    @unittest.skipUnless(importlib.util.find_spec("docx"), "python-docx is not installed")
    def test_extracts_docx_paragraphs_and_tables(self):
        from docx import Document
        from io import BytesIO

        document = Document()
        document.add_paragraph("Chapter overview")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Term"
        table.cell(0, 1).text = "Meaning"
        stream = BytesIO()
        document.save(stream)
        pages = extract_document("notes.docx", stream.getvalue())
        self.assertIn("Chapter overview", pages[0].text)
        self.assertIn("Term | Meaning", pages[0].text)

    @unittest.skipUnless(importlib.util.find_spec("docx"), "python-docx is not installed")
    def test_rejects_corrupt_and_empty_docx_files(self):
        from docx import Document
        from io import BytesIO

        with self.assertRaises(DocumentError):
            extract_document("corrupt.docx", b"not a DOCX")
        document = Document()
        stream = BytesIO()
        document.save(stream)
        with self.assertRaises(DocumentError):
            extract_document("empty.docx", stream.getvalue())


if __name__ == "__main__":
    unittest.main()
